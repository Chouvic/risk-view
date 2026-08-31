"""Format-specific readers: CSV and Excel in, uniform {header: raw string} rows out.

Every reader returns the same shape, so validation is identical regardless of source format.
"""

import csv
import io
from datetime import datetime

from openpyxl import load_workbook


class UnsupportedFormatError(ValueError):
    pass


def read_rows(data: bytes, filename: str) -> list[dict[str, str]]:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    reader = _READERS.get(extension)
    if reader is None:
        raise UnsupportedFormatError(
            f"unsupported file type {extension!r}; expected one of {sorted(_READERS)}"
        )
    return reader(data)


def _read_csv(data: bytes) -> list[dict[str, str]]:
    stream = io.StringIO(data.decode("utf-8-sig"))
    return [
        {header: (value or "") for header, value in record.items()}
        for record in csv.DictReader(stream)
    ]


def _read_xlsx(data: bytes) -> list[dict[str, str]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    try:
        headers = [_as_text(cell) for cell in next(rows)]
    except StopIteration:
        return []
    return [dict(zip(headers, (_as_text(cell) for cell in row))) for row in rows]


def _as_text(cell: object) -> str:
    """Render a spreadsheet cell the way it would appear in a CSV export."""
    if cell is None:
        return ""
    if isinstance(cell, datetime):
        return cell.isoformat(sep=" ")
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell)


_READERS = {"csv": _read_csv, "xlsx": _read_xlsx}
